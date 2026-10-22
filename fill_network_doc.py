import os
import docx
from docx import Document
from docx.shared import Inches

template_path = r"C:/Users/hp/Downloads/Lembar_Jawaban_UTS_Pemrograman_Jaringan_Proyek.docx"
output_docx = r"C:/Users/hp/UTS-Pemrograman-Jaringan/Lembar_Jawaban_UTS_Pemrograman_Jaringan_240104013_Yusuf_Rizal_Indrianto.docx"
img_dir = r"C:/Users/hp/UTS-Pemrograman-Jaringan/images"

doc = Document(template_path)

def set_cell(table, r, c, text):
    table.rows[r].cells[c].text = text

# -------------------------------------------------------------
# IDENTITAS MAHASISWA & PROYEK (Table 0)
# -------------------------------------------------------------
set_cell(doc.tables[0], 0, 1, "Yusuf Rizal Indrianto")
set_cell(doc.tables[0], 1, 1, "240104013")
set_cell(doc.tables[0], 2, 1, "TK24A1 / Semester V")
set_cell(doc.tables[0], 3, 1, "Sistem Monitoring & Kontrol Perangkat Jaringan Laboratorium Berbasis Multi-Protokol Socket Python (SIP-LAB)")
set_cell(doc.tables[0], 4, 1, "21 Oktober 2026 (Topik: Lab Network & Device Monitoring System)")
set_cell(doc.tables[0], 5, 1, "24 Oktober 2026")

# -------------------------------------------------------------
# RINGKASAN KEUNIKAN PROYEK (Table 1)
# -------------------------------------------------------------
set_cell(doc.tables[1], 0, 1, "Laboran kesulitan memantau puluhan PC dan switch laboratorium secara manual. Aplikasi ini menyediakan sentralisasi kontrol status perangkat, telemetri suhu & daya, transfer konfigurasi, dan dashboard web via socket.")
set_cell(doc.tables[1], 1, 1, "CHECK_DEVICE <ID> : Mengambil metrik performa real-time perangkat (status online/offline, beban CPU, persentase RAM, dan lokasi meja lab).")
set_cell(doc.tables[1], 2, 1, "REBOOT_DEVICE <ID> : Mengirimkan sinyal instruksi restart perangkat jaringan/PC lab dari jarak jauh dan mereset beban pemakaian CPU/RAM.")
set_cell(doc.tables[1], 3, 1, "GET_TELEMETRY (UDP) : Layanan broadcast polling telemetri kondisi fisik ruangan lab (suhu ruang Celsius, kelembapan udara %, dan konsumsi daya listrik Watt).")
set_cell(doc.tables[1], 4, 1, "Berkas konfigurasi jaringan (lab_network_config.ini) dan berkas rekaman log audit keamanan perangkat lab (lab_audit_event.log).")
set_cell(doc.tables[1], 5, 1, "Rute '/' menampilkan dashboard antarmuka HTML daftar perangkat, sedangkan rute '/status' menyajikan data telemetri JSON secara real-time.")

# -------------------------------------------------------------
# TAHAP 1: TOPIK DAN LINGKUNGAN PYTHON
# -------------------------------------------------------------
# Table 2: 1.1 Tujuan dan rancangan komunikasi
set_cell(doc.tables[2], 0, 1, "SIP-LAB adalah sistem pemantauan dan kontrol terpadu perangkat laboratorium komputer menggunakan socket Python. Aplikasi ini memungkinkan staf admin memonitor status kesehatan perangkat, mengirim instruksi remote reboot, membaca telemetri lingkungan, mentransfer file konfigurasi, serta menyediakan pemantauan web.")
set_cell(doc.tables[2], 1, 1, "Klien mengirimkan instruksi teks (CHECK_DEVICE, REBOOT_DEVICE, LIST_DEVICES, QUIT) via TCP, meminta telemetri sensor via datagram UDP, mentransfer file biner, dan mengakses halaman via HTTP client.")
set_cell(doc.tables[2], 2, 1, "Server memvalidasi instruksi, mengolah basis data status perangkat, menghitung data telemetri, menyimpan berkas biner secara utuh, dan menyajikan payload respons HTML/JSON.")

# Table 3: 1.2 Pemeriksaan lingkungan dan pengaturan jaringan
set_cell(doc.tables[3], 0, 2, "Python 3.11.9 (64-bit)")
set_cell(doc.tables[3], 1, 2, "Siap (Loopback 127.0.0.1 aktif, latency < 1 ms, 0% packet loss)")
set_cell(doc.tables[3], 2, 1, "127.0.0.1 : Port 5000 (TCP Dasar) & Port 5050 (TCP Multiklien)")
set_cell(doc.tables[3], 2, 2, "Siap (Port berhasil dibuka dan ditutup)")
set_cell(doc.tables[3], 3, 1, "127.0.0.1 : Port 6000 (UDP Telemetri)")
set_cell(doc.tables[3], 3, 2, "Siap (Socket datagram UDP beroperasi normal)")
set_cell(doc.tables[3], 4, 1, "127.0.0.1 : Port 8080 (HTTP Server Socket)")
set_cell(doc.tables[3], 4, 2, "Siap (Merespons rute / dan /status)")

# Table 4: 1.3 Struktur folder proyek
set_cell(doc.tables[4], 0, 1, "Folder: UTS-Pemrograman-Jaringan (Lokasi: C:\\Users\\hp\\UTS-Pemrograman-Jaringan)")
set_cell(doc.tables[4], 1, 1, "- tcp_server.py & tcp_client.py (Tahap 2)\n- udp_server.py & udp_client.py (Tahap 3)\n- tcp_multiclient_server.py (Tahap 4)\n- file_transfer_server.py & file_transfer_client.py (Tahap 5)\n- http_server.py & http_client.py (Tahap 6)\n- test_all.py (Otomasi Pengujian)\n- README.md & Berkas Data Uji (.ini & .log)")

# Table 5: Bukti Tahap 1 (Image)
cell_t5 = doc.tables[5].rows[1].cells[0]
cell_t5.text = "Hasil pemeriksaan lingkungan terminal menunjukkan Python 3.11.9 terpasang, konektivitas loopback 127.0.0.1 aktif tanpa loss, dan seluruh alokasi port siap digunakan:"
p5 = cell_t5.add_paragraph()
p5.add_run().add_picture(os.path.join(img_dir, "bukti_tahap1_lingkungan.png"), width=Inches(5.7))

# -------------------------------------------------------------
# TAHAP 2: KOMUNIKASI TCP CLIENT-SERVER
# -------------------------------------------------------------
# Table 6: 2.1 Dua perintah utama
set_cell(doc.tables[6], 0, 0, "CHECK_DEVICE")
set_cell(doc.tables[6], 0, 1, "CHECK_DEVICE <DEVICE_ID>\nContoh: CHECK_DEVICE PC-LAB-01")
set_cell(doc.tables[6], 0, 2, "Server mencocokkan ID dengan dictionary 'devices', mengekstrak status online/offline, persentase CPU, RAM, dan lokasi meja lab.")
set_cell(doc.tables[6], 0, 3, "OK: PC-LAB-01 Status=ONLINE, CPU=45%, RAM=62%, Lokasi=Lab 4 Meja 1")

set_cell(doc.tables[6], 1, 0, "REBOOT_DEVICE")
set_cell(doc.tables[6], 1, 1, "REBOOT_DEVICE <DEVICE_ID>\nContoh: REBOOT_DEVICE PC-LAB-02")
set_cell(doc.tables[6], 1, 2, "Server memverifikasi status perangkat. Jika online, server mengirim sinyal restart virtual dan mereset metrik CPU menjadi 12% dan RAM menjadi 28%.")
set_cell(doc.tables[6], 1, 3, "OK: Sinyal reboot berhasil dikirim ke PC-LAB-02. Sistem telah di-reset (CPU=12%, RAM=28%).")

# Table 7: 2.2 Implementasi dan alur komunikasi
set_cell(doc.tables[7], 0, 1, "tcp_server.py (C:\\Users\\hp\\UTS-Pemrograman-Jaringan\\tcp_server.py)")
set_cell(doc.tables[7], 1, 1, "tcp_client.py (C:\\Users\\hp\\UTS-Pemrograman-Jaringan\\tcp_client.py)")
set_cell(doc.tables[7], 2, 1, "Server menjalankan socket() -> setsockopt() -> bind() ke 127.0.0.1:5000 -> listen(). Saat klien memanggil connect(), server menjalankan accept() menghasilkan socket baru untuk klien. Klien dan server saling bertukar data menggunakan sendall() dan recv(1024).")
set_cell(doc.tables[7], 3, 1, "Klien mengirim perintah 'QUIT'. Server mendeteksi perintah tersebut, mengirim balasan penutup 'BYE: Menutup sesi koneksi TCP', lalu menutup socket klien via close(). Server tetap melanjutkan perulangan while True (listening) menunggu klien berikutnya.")

# Table 8: 2.3 Pengujian
set_cell(doc.tables[8], 0, 1, "CHECK_DEVICE PC-LAB-01")
set_cell(doc.tables[8], 0, 2, "OK: PC-LAB-01 Status=ONLINE, CPU=45%, RAM=62%, Lokasi=Lab 4 Meja 1")
set_cell(doc.tables[8], 0, 3, "Berhasil")

set_cell(doc.tables[8], 1, 1, "REBOOT_DEVICE PC-LAB-02")
set_cell(doc.tables[8], 1, 2, "OK: Sinyal reboot berhasil dikirim ke PC-LAB-02. Sistem telah di-reset (CPU=12%, RAM=28%).")
set_cell(doc.tables[8], 1, 3, "Berhasil")

set_cell(doc.tables[8], 2, 1, "QUIT")
set_cell(doc.tables[8], 2, 2, "BYE: Menutup sesi koneksi TCP. Terima kasih. (Server tetap aktif listening)")
set_cell(doc.tables[8], 2, 3, "Berhasil")

# Table 9: Bukti Tahap 2 (Image)
cell_t9 = doc.tables[9].rows[1].cells[0]
cell_t9.text = "Bukti eksekusi komunikasi TCP: server memproses perintah khusus CHECK_DEVICE dan REBOOT_DEVICE, serta penutupan sesi secara tertib tanpa mematikan server:"
p9 = cell_t9.add_paragraph()
p9.add_run().add_picture(os.path.join(img_dir, "bukti_tahap2_tcp.png"), width=Inches(5.7))

# -------------------------------------------------------------
# TAHAP 3: KOMUNIKASI UDP DAN TIMEOUT
# -------------------------------------------------------------
# Table 10: 3.1 Fungsi layanan pendukung
set_cell(doc.tables[10], 0, 1, "Layanan Telemetri Sensor Lingkungan Lab (suhu ruang, kelembapan, dan daya listrik) yang membutuhkan pengiriman data berkala cepat tanpa overhead handshaking koneksi.")
set_cell(doc.tables[10], 1, 1, "Server: udp_server.py  |  Klien: udp_client.py")
set_cell(doc.tables[10], 2, 1, "Klien mengirim pesan datagram teks: 'GET_TELEMETRY' via sendto().")
set_cell(doc.tables[10], 3, 1, "TELEMETRY: Temp=24.1C, Humidity=55%, PowerUsage=1240W, Timestamp=17:15:22 via recvfrom().")
set_cell(doc.tables[10], 4, 1, "Durasi timeout diatur sebesar 3.0 detik (sock.settimeout(3.0)). Apabila server tidak membalas dalam 3 detik, klien menangkap exception socket.timeout, menampilkan pesan error, dan tidak mengalami freeze/hang.")

# Table 11: 3.2 Pengujian UDP
set_cell(doc.tables[11], 0, 1, "Kirim datagram 'GET_TELEMETRY' saat udp_server.py aktif")
set_cell(doc.tables[11], 0, 2, "Menerima balasan telemetri lengkap: Temp=24.1C, Humidity=55%, PowerUsage=1240W")
set_cell(doc.tables[11], 0, 3, "Berhasil")

set_cell(doc.tables[11], 1, 1, "Kirim datagram 'GET_TELEMETRY' saat udp_server.py TIDAK aktif")
set_cell(doc.tables[11], 1, 2, "Muncul notifikasi: '[!] TIMEOUT: Batas waktu 3.0 detik terlewati. Server UDP 127.0.0.1:6000 tidak merespons.'")
set_cell(doc.tables[11], 1, 3, "Berhasil")

# Table 12: 3.3 Alasan penggunaan protokol
set_cell(doc.tables[12], 0, 1, "TCP bersifat connection-oriented dan reliable (menjamin pengiriman urut & bebas galat). Sangat krusial untuk perintah kontrol perangkat seperti REBOOT_DEVICE dan pengecekan status di mana instruksi tidak boleh hilang atau rusak.")
set_cell(doc.tables[12], 1, 1, "UDP bersifat connectionless, ringan (low overhead), dan berkecepatan tinggi. Sangat ideal untuk telemetri sensor periodik (suhu/daya) di mana kehilangan 1 paket datagram tidak merusak sistem karena akan segera diperbarui oleh sampel berikutnya.")

# Table 13: Bukti Tahap 3 (Image)
cell_t13 = doc.tables[13].rows[1].cells[0]
cell_t13.text = "Bukti pengujian layanan UDP saat server aktif (respons normal) dan saat server dimatikan (penanganan timeout 3.0 detik):"
p13 = cell_t13.add_paragraph()
p13.add_run().add_picture(os.path.join(img_dir, "bukti_tahap3_udp.png"), width=Inches(5.7))

# -------------------------------------------------------------
# TAHAP 4: SERVER TCP MULTIKLIEN
# -------------------------------------------------------------
# Table 14: 4.1 Penerapan threading
set_cell(doc.tables[14], 0, 1, "tcp_multiclient_server.py (C:\\Users\\hp\\UTS-Pemrograman-Jaringan\\tcp_multiclient_server.py)")
set_cell(doc.tables[14], 1, 1, "Thread utama menjalankan loop server.accept(). Ketika ada klien baru yang terhubung, thread utama tidak melayani sendiri, melainkan langsung membuat dan memulai objek threading.Thread baru, lalu segera kembali menunggu klien berikutnya.")
set_cell(doc.tables[14], 2, 1, "Setiap thread pekerja (worker thread) mengeksekusi fungsi handle_client() khusus untuk satu socket klien: menyapa klien, memproses perintah dalam perulangan independen, dan menutup koneksi secara terisolasi tanpa memblokir klien lain.")
set_cell(doc.tables[14], 3, 1, "Nama thread (threading.current_thread().name, contoh: Thread-Klien-1), alamat IP & port ephemeral klien, serta jumlah thread yang sedang aktif.")

# Table 15: 4.2 Pengujian tiga klien
set_cell(doc.tables[15], 0, 1, "CHECK_DEVICE PC-LAB-01")
set_cell(doc.tables[15], 0, 2, "Thread-Klien-1 (Klien 1 Meja 1)")
set_cell(doc.tables[15], 0, 3, "OK: PC-LAB-01 Status=ONLINE, CPU=45%, RAM=62%, Lokasi=Lab 4 Meja 1")

set_cell(doc.tables[15], 1, 1, "REBOOT_DEVICE PC-LAB-02")
set_cell(doc.tables[15], 1, 2, "Thread-Klien-2 (Klien 2 Meja 2)")
set_cell(doc.tables[15], 1, 3, "OK: Sinyal reboot berhasil dikirim ke PC-LAB-02. Sistem telah di-reset (CPU=12%, RAM=28%).")

set_cell(doc.tables[15], 2, 1, "LIST_DEVICES")
set_cell(doc.tables[15], 2, 2, "Thread-Klien-3 (Klien 3 Admin Lab)")
set_cell(doc.tables[15], 2, 3, "OK: Daftar Perangkat Lab: [PC-LAB-01, PC-LAB-02, PC-LAB-03, SW-CORE-01]")

# Table 16: Diagram alur multiclient (Image)
cell_t16 = doc.tables[16].rows[1].cells[0]
cell_t16.text = "Diagram arsitektur alur server multiklien: thread utama menerima koneksi dan mendelegasikan pemrosesan ke thread pekerja independen untuk 3 klien simultan:"
p16 = cell_t16.add_paragraph()
p16.add_run().add_picture(os.path.join(img_dir, "diagram_tahap4_multiclient.png"), width=Inches(5.7))

# -------------------------------------------------------------
# TAHAP 5: TRANSFER BERKAS MELALUI TCP
# -------------------------------------------------------------
# Table 18: 5.2 Implementasi pengiriman
set_cell(doc.tables[18], 0, 1, "file_transfer_server.py & file_transfer_client.py")
set_cell(doc.tables[18], 1, 1, "Panjang nama berkas dienkode 2 byte (maks. 65.535 karakter). Ukuran berkas dienkode 8 byte (maks. 18 Exabyte). Program memvalidasi nama berkas tidak boleh kosong dan ukuran harus > 0.")
set_cell(doc.tables[18], 2, 1, "Server menggunakan fungsi loop 'recv_exact()' dan buffer chunking 4096 byte yang terus mengumpulkan data ke variabel bytearray hingga genap mencapai nilai filesize dari header biner.")
set_cell(doc.tables[18], 3, 1, "1. lab_network_config.ini (berkas konfigurasi IP & ambang batas alert lab)\n2. lab_audit_event.log (berkas riwayat log aktivitas dan insiden perangkat lab)")

# Table 19: 5.3 Pengujian dua berkas
set_cell(doc.tables[19], 0, 0, "lab_network_config.ini")
set_cell(doc.tables[19], 0, 1, "445 byte")
set_cell(doc.tables[19], 0, 2, "445 byte")
set_cell(doc.tables[19], 0, 3, "Ukuran sama persis (100% identik). Isi konfigurasi jaringan tersimpan utuh di folder received_files.")

set_cell(doc.tables[19], 1, 0, "lab_audit_event.log")
set_cell(doc.tables[19], 1, 1, "570 byte")
set_cell(doc.tables[19], 1, 2, "570 byte")
set_cell(doc.tables[19], 1, 3, "Ukuran sama persis (100% identik). Seluruh baris log rekaman audit tersimpan utuh tanpa korupsi data.")

# Table 20: Bukti Tahap 5 (Image)
cell_t20 = doc.tables[20].rows[1].cells[0]
cell_t20.text = "Bukti proses pengiriman berkas dengan protokol header biner [2B Nama Len][Nama][8B Size][Data] serta perbandingan ukuran berkas asli dan diterima:"
p20 = cell_t20.add_paragraph()
p20.add_run().add_picture(os.path.join(img_dir, "bukti_tahap5_file_transfer.png"), width=Inches(5.7))

# -------------------------------------------------------------
# TAHAP 6: HTTP SERVER DAN KLIEN
# -------------------------------------------------------------
# Table 21: 6.1 Layanan HTTP sesuai topik proyek
set_cell(doc.tables[21], 0, 1, "HTML")
set_cell(doc.tables[21], 0, 2, "Halaman web antarmuka pemantauan Lab Network Monitoring (daftar perangkat PC/Switch, lokasi meja, badge status online/offline, persentase CPU & RAM).")

set_cell(doc.tables[21], 1, 1, "JSON")
set_cell(doc.tables[21], 1, 2, "Struktur JSON data telemetri suhu (24.5C), kelembapan (58%), konsumsi daya (1280W), dan array status seluruh perangkat lab.")

set_cell(doc.tables[21], 2, 0, "/rute_tidak_ada (404)")
set_cell(doc.tables[21], 2, 1, "HTML Pesan Kesalahan")
set_cell(doc.tables[21], 2, 2, "Halaman HTML error '404 Not Found' yang menginformasikan bahwa sumber daya pada URL yang diminta tidak tersedia di server.")

# Table 22: 6.2 Implementasi dan struktur respons
set_cell(doc.tables[22], 0, 1, "http_server.py (C:\\Users\\hp\\UTS-Pemrograman-Jaringan\\http_server.py)")
set_cell(doc.tables[22], 1, 1, "http_client.py (C:\\Users\\hp\\UTS-Pemrograman-Jaringan\\http_client.py)")
set_cell(doc.tables[22], 2, 1, "GET / HTTP/1.1\\r\\nHost: 127.0.0.1:8080\\r\\nUser-Agent: RawSocketPythonClient/1.0\\r\\nConnection: close\\r\\n\\r\\n")
set_cell(doc.tables[22], 3, 1, "Status Line: HTTP/1.1 200 OK | Headers: Server: LabNetMonitor/1.0, Content-Type: text/html (atau application/json), Content-Length: [N], Connection: close")
set_cell(doc.tables[22], 4, 1, "String HTML dengan tabel monitoring berformat styling CSS modern, atau string JSON terenkode UTF-8 berisi data telemetri lab.")

# Table 23: 6.3 Pengujian
set_cell(doc.tables[23], 0, 1, "HTTP/1.1 200 OK")
set_cell(doc.tables[23], 0, 2, "Halaman HTML antarmuka dashboard monitoring berhasil ditampilkan utuh di klien.")

set_cell(doc.tables[23], 1, 1, "HTTP/1.1 200 OK")
set_cell(doc.tables[23], 1, 2, "Data JSON berhasil diterima dan di-parse dengan Content-Type: application/json.")

set_cell(doc.tables[23], 2, 0, "/rute_tidak_ada")
set_cell(doc.tables[23], 2, 1, "HTTP/1.1 404 Not Found")
set_cell(doc.tables[23], 2, 2, "Muncul pesan respons 404 bahwa URL tidak ditemukan.")

# Table 24: Bukti Tahap 6 (Image)
cell_t24 = doc.tables[24].rows[1].cells[0]
cell_t24.text = "Bukti pengujian server dan klien HTTP socket murni untuk respons 200 OK (HTML & JSON) serta respons 404 Not Found:"
p24 = cell_t24.add_paragraph()
p24.add_run().add_picture(os.path.join(img_dir, "bukti_tahap6_http.png"), width=Inches(5.7))

# -------------------------------------------------------------
# RINGKASAN AKHIR DAN PENGUMPULAN (Table 25)
# -------------------------------------------------------------
set_cell(doc.tables[25], 0, 1, "Pemeriksaan Lingkungan & Socket Test")
set_cell(doc.tables[25], 0, 2, "Memastikan lingkungan Python 3.11 dan seluruh port (5000, 5050, 6000, 7000, 8080) siap digunakan.")

set_cell(doc.tables[25], 1, 1, "TCP Dasar: CHECK_DEVICE & REBOOT_DEVICE")
set_cell(doc.tables[25], 1, 2, "Menyediakan jalur kontrol handal untuk mengecek status kesehatan perangkat lab dan merestart jarak jauh.")

set_cell(doc.tables[25], 2, 1, "UDP Telemetri Sensor & Timeout")
set_cell(doc.tables[25], 2, 2, "Menyediakan layanan monitoring cepat kondisi fisik ruang lab (suhu, kelembapan, daya listrik) tanpa overhead.")

set_cell(doc.tables[25], 3, 1, "TCP Multiklien Threading")
set_cell(doc.tables[25], 3, 2, "Memungkinkan banyak asisten lab atau teknisi mengakses dan mengontrol perangkat lab secara bersamaan.")

set_cell(doc.tables[25], 4, 1, "Transfer Berkas Protokol Biner")
set_cell(doc.tables[25], 4, 2, "Mengirimkan berkas konfigurasi jaringan (.ini) dan audit log (.log) secara utuh antar-komputer lab.")

set_cell(doc.tables[25], 5, 1, "HTTP Server & Client Raw Socket")
set_cell(doc.tables[25], 5, 2, "Menyajikan dashboard web HTML dan API data JSON yang dapat diakses melalui browser oleh pimpinan lab.")

doc.save(output_docx)
print(f"Berhasil mengisi dan menyimpan dokumen di: {output_docx}")
